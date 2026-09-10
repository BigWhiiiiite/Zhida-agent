"""Run with: .venv/bin/python smoke_test.py"""
import asyncio
import logging
import os
from pathlib import Path
from tempfile import TemporaryDirectory

import httpx
from agents.usage import Usage
from docx import Document
from fastapi.testclient import TestClient
from openai import APIConnectionError

from app import main, storage
from app.agent import _profile_from_model_text
from app.browser_models import BrowserSnapshot, PageField
from app.form_agent import _form_plan_from_model_text, _local_safe_plan
from app.extractors import extract_text
from app.model_provider import normalize_proxy_response
from app.models import CandidateProfile, Education, ModelHealth
from app.job_recommendations import recommendation_batch

os.environ["APP_AGENT_MODE"] = "rules"


async def verify_proxy_normalization() -> None:
    response = httpx.Response(
        200,
        headers={"content-type": "application/json"},
        request=httpx.Request("POST", "https://proxy.example/v1/responses"),
        json={
            "usage": {
                "input_tokens": 10,
                "input_tokens_details": {"cached_tokens": 0},
                "output_tokens": 5,
                "output_tokens_details": {"reasoning_tokens": 0},
                "total_tokens": 15,
            }
        },
    )
    await normalize_proxy_response(response)
    assert response.json()["usage"]["input_tokens_details"]["cache_write_tokens"] == 0


asyncio.run(verify_proxy_normalization())
assert Usage().input_tokens_details.cached_tokens == 0
assert logging.getLogger("openai.agents").level == logging.CRITICAL
fallback_profile = _profile_from_model_text(
    '```json\n{"name":"备用模型测试","email":"fallback@example.com"}\n```'
)
assert fallback_profile.name == "备用模型测试"
assert fallback_profile.email == "fallback@example.com"

fallback_plan = _form_plan_from_model_text(
    '```json\n{"page_summary":"test","site_type":"lever","actions":[],"missing_questions":[]}\n```'
)
assert fallback_plan.site_type == "lever"

local_plan = _local_safe_plan(
    BrowserSnapshot(session_id="test", url="https://jobs.lever.co/example/apply", title="Test", fields=[
        PageField(selector="[data-zhida-field=name]", label="Full name", name="name", required=True),
        PageField(selector="[data-zhida-field=email]", label="Email", name="email", required=True),
        PageField(selector="[data-zhida-field=gender]", label="Gender", name="gender", required=True),
        PageField(selector="[data-zhida-field=resume]", label="Resume/CV", name="resume", field_type="file", required=True),
    ]),
    CandidateProfile(name="Test Candidate", email="test@example.com"),
)
assert [action.action for action in local_plan.actions] == ["fill", "fill", "ask_user", "skip"]
assert local_plan.actions[2].sensitive

relationship_plan = _local_safe_plan(
    BrowserSnapshot(session_id="relationship", url="https://careers.example/apply", title="Test", fields=[
        PageField(selector="[data-zhida-field=emergency-name]", label="紧急联系人姓名", required=True),
        PageField(selector="[data-zhida-field=emergency-phone]", label="紧急联系人电话", required=True),
    ]),
    CandidateProfile(name="候选人本人", phone="13800138000"),
)
assert [action.action for action in relationship_plan.actions] == ["ask_user", "ask_user"]
assert all(action.sensitive and not action.value for action in relationship_plan.actions)

decision_fields = [
    PageField(selector="[data-zhida-field=transfer-yes]", label="是否接受调剂 — 是", name="transfer",
              field_type="radio", group_label="是否接受调剂", option_label="是", option_value="yes",
              options=["是", "否"], required=True),
    PageField(selector="[data-zhida-field=transfer-no]", label="是否接受调剂 — 否", name="transfer",
              field_type="radio", group_label="是否接受调剂", option_label="否", option_value="no",
              options=["是", "否"], required=True),
]
decision_plan = _local_safe_plan(
    BrowserSnapshot(session_id="decision", url="https://careers.example/apply", title="Test",
                    fields=decision_fields),
    CandidateProfile(),
)
assert [action.action for action in decision_plan.actions] == ["ask_user", "ask_user"]
remembered_decision = _local_safe_plan(
    BrowserSnapshot(session_id="decision", url="https://careers.example/apply", title="Test",
                    fields=decision_fields),
    CandidateProfile(application_answers={"是否接受调剂": "否"}),
)
assert [action.action for action in remembered_decision.actions] == ["skip", "check"]
assert remembered_decision.actions[1].user_confirmed

education_plan = _local_safe_plan(
    BrowserSnapshot(session_id="education", url="https://careers.example/apply", title="Test", fields=[
        PageField(selector="[data-zhida-field=level]", label="最高学历", field_type="select-one",
                  options=["高中", "大学专科", "大学本科", "硕士研究生"]),
        PageField(selector="[data-zhida-field=degree]", label="学位", field_type="select-one",
                  options=["学士", "硕士", "博士"]),
    ]),
    CandidateProfile(education=[Education(school="Test University", degree="本科")]),
)
assert [action.value for action in education_plan.actions] == ["大学本科", "学士"]

learned_plan = _local_safe_plan(
    BrowserSnapshot(session_id="learned", url="https://careers.example/apply", title="Test", fields=[
        PageField(selector="[data-zhida-field=qq]", label="QQ号", name="qq", required=True),
        PageField(selector="[data-zhida-field=community]", label="你最常参与的技术社区", name="community"),
    ]),
    CandidateProfile(qq="12345678", application_answers={"你最常参与的技术社区": "GitHub"}),
)
assert [action.value for action in learned_plan.actions] == ["12345678", "GitHub"]
assert learned_plan.actions[0].value_source == "主档案.qq"
assert learned_plan.actions[1].value_source.startswith("主档案.application_answers")

graduating_profile = CandidateProfile(
    target_role="Agent 开发工程师", skills=["Python", "FastAPI", "Agent"],
    education=[Education(school="Test University", end_date="2026.12")],
)
graduating_jobs = recommendation_batch(graduating_profile).jobs
baidu_agent = next(item for item in graduating_jobs if item.job.job_code == "J101017")
assert baidu_agent.graduation_match is True

shenzhen_batch = recommendation_batch(graduating_profile, "深圳市")
assert shenzhen_batch.selected_location == "深圳"
assert "深圳" in shenzhen_batch.available_locations
assert shenzhen_batch.jobs
assert all(item.location_match is True for item in shenzhen_batch.jobs)
assert all(any("深圳" in location for location in item.job.locations) for item in shenzhen_batch.jobs)
assert all(any("所选城市：深圳" in reason for reason in item.reasons) for item in shenzhen_batch.jobs)


class UnavailableExtractor:
    name = "unavailable-test-provider"

    async def parse(self, _: str):
        raise APIConnectionError(request=httpx.Request("POST", "https://proxy.example/v1/responses"))


with TemporaryDirectory() as temporary:
    root = Path(temporary)
    merged_docx = root / "merged.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "重复标题"
    table.cell(0, 1).text = "第一行"
    table.cell(1, 1).text = "第二行"
    document.save(merged_docx)
    assert extract_text(merged_docx).count("重复标题") == 1

    storage.DB_PATH = root / "test.db"
    main.UPLOAD_DIR = root / "uploads"

    with TestClient(main.app) as client:
        assert client.get("/api/health").json()["product"] == "Zhida"
        assert client.get("/api/profile").status_code == 401
        registered = client.post("/api/auth/register", json={
            "email": "alice@example.com", "password": "Alice-pass-2026", "display_name": "Alice",
        })
        assert registered.status_code == 201, registered.text
        assert registered.json()["user"]["email"] == "alice@example.com"
        assert "password" not in registered.text.lower()
        assert "httponly" in registered.headers["set-cookie"].lower()
        assert client.get("/api/auth/me").json()["display_name"] == "Alice"
        original_health_check = main.check_model_health
        main.check_model_health = lambda: asyncio.sleep(0, result=ModelHealth(
            status="ok", model="test-model", latency_ms=12, message="主模型当前可用"
        ))
        try:
            model_health = client.post("/api/model/health")
            assert model_health.status_code == 200
            assert model_health.json()["status"] == "ok"
            assert "key" not in model_health.text.lower()
        finally:
            main.check_model_health = original_health_check
        initial = client.get("/api/profile")
        assert initial.status_code == 200

        remembered_qq = client.post("/api/profile/application-answer", json={
            "question": "QQ号", "field_name": "candidate_qq", "value": "12345678",
        })
        assert remembered_qq.status_code == 200
        assert remembered_qq.json()["qq"] == "12345678"
        remembered_custom = client.post("/api/profile/application-answer", json={
            "question": "你最常参与的技术社区", "field_name": "community", "value": "GitHub",
        })
        assert remembered_custom.status_code == 200
        assert remembered_custom.json()["application_answers"]["你最常参与的技术社区"] == "GitHub"
        rejected_sensitive = client.post("/api/profile/application-answer", json={
            "question": "是否同意隐私条款", "field_name": "consent", "value": "是",
        })
        assert rejected_sensitive.status_code == 422

        first_text = (
            "姓名：李春博\n性别：男\n年龄：24\n邮箱：first@example.com\n"
            "手机号：13800138000\n技能\nPython，FastAPI\n项目经历\nAgent 求职助手"
        ).encode()
        first = client.post("/api/resumes", files={"file": ("agent.txt", first_text, "text/plain")})
        assert first.status_code == 201, first.text
        first_json = first.json()
        assert first_json["profile"]["name"] == "李春博"
        assert first_json["evidence"]

        assert client.post("/api/auth/logout").status_code == 204
        bob = client.post("/api/auth/register", json={
            "email": "bob@example.com", "password": "Bob-pass-2026", "display_name": "Bob",
        })
        assert bob.status_code == 201, bob.text
        assert client.get("/api/resumes").json() == []
        assert client.get("/api/profile").json()["email"] == ""
        assert client.get(f"/api/resumes/{first_json['id']}").status_code == 404
        assert client.post("/api/auth/logout").status_code == 204
        assert client.post("/api/auth/login", json={
            "email": "alice@example.com", "password": "incorrect-password",
        }).status_code == 401
        signed_in = client.post("/api/auth/login", json={
            "email": "ALICE@example.com", "password": "Alice-pass-2026",
        })
        assert signed_in.status_code == 200, signed_in.text
        assert client.get("/api/resumes").json()[0]["id"] == first_json["id"]

        duplicate = client.post("/api/resumes", files={"file": ("copy.txt", first_text, "text/plain")})
        assert duplicate.status_code == 409

        profile = client.get("/api/profile").json()
        assert profile["email"] == "first@example.com"

        recommendations = client.get("/api/jobs/recommendations")
        assert recommendations.status_code == 200
        recommendation_json = recommendations.json()
        assert recommendation_json["engine"] == "local-explainable-v1"
        assert len(recommendation_json["jobs"]) >= 5
        scores = [item["match_score"] for item in recommendation_json["jobs"]]
        assert scores == sorted(scores, reverse=True)
        assert all(item["reasons"] for item in recommendation_json["jobs"])
        assert all(item["job"]["source_url"].startswith("https://") for item in recommendation_json["jobs"])

        shenzhen_recommendations = client.get("/api/jobs/recommendations", params={"location": "深圳"})
        assert shenzhen_recommendations.status_code == 200
        shenzhen_json = shenzhen_recommendations.json()
        assert shenzhen_json["selected_location"] == "深圳"
        assert shenzhen_json["jobs"]
        assert all("深圳" in item["job"]["locations"] for item in shenzhen_json["jobs"])

        first_job_id = recommendation_json["jobs"][0]["job"]["id"]
        queued = client.post("/api/jobs/queue", json={"job_ids": [first_job_id], "resume_id": first_json["id"]})
        assert queued.status_code == 201, queued.text
        assert queued.json()[0]["job_id"] == first_job_id
        assert queued.json()[0]["resume_id"] == first_json["id"]
        queue_id = queued.json()[0]["id"]
        assert client.get("/api/jobs/queue").json()[0]["recommendation"]["reasons"]
        assert client.delete(f"/api/jobs/queue/{queue_id}").status_code == 204
        assert client.get("/api/jobs/queue").json() == []
        unknown_job = client.post("/api/jobs/queue", json={"job_ids": ["not-a-job"], "resume_id": ""})
        assert unknown_job.status_code == 422

        second_text = "姓名：李春博\n邮箱：new@example.com\n技能\nPydanticAI，React".encode()
        second = client.post("/api/resumes", files={"file": ("english.txt", second_text, "text/plain")})
        assert second.status_code == 201, second.text
        conflicts = client.get("/api/conflicts").json()
        email_conflict = next(item for item in conflicts if item["field_path"] == "email")
        resolved = client.post(f"/api/conflicts/{email_conflict['id']}/resolve", json={"choice": "incoming"})
        assert resolved.status_code == 200
        assert client.get("/api/profile").json()["email"] == "new@example.com"

        evidence_id = first_json["evidence"][0]["id"]
        reviewed = client.patch(
            f"/api/resumes/{first_json['id']}/evidence/{evidence_id}", json={"status": "confirmed"}
        )
        assert reviewed.status_code == 200
        assert reviewed.json()["evidence"][0]["status"] == "confirmed"

        assert client.get(f"/api/resumes/{first_json['id']}/download").status_code == 200
        preview = client.get(f"/api/resumes/{first_json['id']}/preview")
        assert preview.status_code == 200
        assert "text/html" in preview.headers["content-type"]

        skills_evidence = next(item for item in first_json["evidence"] if item["field_path"] == "skills")
        edited = client.patch(
            f"/api/resumes/{first_json['id']}/evidence/{skills_evidence['id']}",
            json={"status": "edited", "value": ["Python", "React"]},
        )
        assert edited.status_code == 200, edited.text
        assert edited.json()["profile"]["skills"] == ["Python", "React"]
        assert client.get("/api/profile").json()["skills"] == ["Python", "React"]

        assert client.post(f"/api/resumes/{first_json['id']}/parse").status_code == 200

        original_extractor = main.get_resume_extractor
        main.get_resume_extractor = lambda: UnavailableExtractor()
        try:
            failed_text = "姓名：网络故障测试\n邮箱：retry@example.com".encode()
            failed = client.post("/api/resumes", files={"file": ("retry.txt", failed_text, "text/plain")})
            assert failed.status_code == 503, failed.text
            failed_id = failed.json()["detail"]["resume_id"]
            retained = client.get(f"/api/resumes/{failed_id}").json()
            assert retained["status"] == "failed"
            assert retained["error_message"]
            assert (main.UPLOAD_DIR / f"{failed_id}.txt").exists()
        finally:
            main.get_resume_extractor = original_extractor
        retried = client.post(f"/api/resumes/{failed_id}/parse")
        assert retried.status_code == 200, retried.text
        assert retried.json()["status"] == "needs_review"

        assert client.get("/api/export").status_code == 200
        assert client.delete(f"/api/resumes/{first_json['id']}").status_code == 204
        assert client.get(f"/api/resumes/{first_json['id']}").status_code == 404

print("Zhida phase-one smoke test passed")
