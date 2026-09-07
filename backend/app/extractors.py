from __future__ import annotations

from html import escape
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


def _docx_blocks(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in their original document order."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _docx_lines(path: Path) -> list[str]:
    document = Document(path)
    lines: list[str] = []
    for block in _docx_blocks(document):
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.extend(line.strip() for line in block.text.splitlines() if line.strip())
            continue

        # python-docx returns the same XML cell more than once for horizontally
        # or vertically merged cells. Track it across the table to avoid the
        # duplicated headings and experience rows seen in extracted resumes.
        seen_cells: set[str] = set()
        for row in block.rows:
            row_values: list[str] = []
            for cell in row.cells:
                cell_key = cell._tc.getroottree().getpath(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                value = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if value:
                    row_values.append(value)
            if row_values:
                lines.extend(line.strip() for line in " | ".join(row_values).splitlines() if line.strip())

    # Layout-heavy resumes sometimes contain adjacent duplicate paragraphs even
    # outside tables. Removing only adjacent duplicates preserves repeated facts
    # that legitimately appear in different sections.
    cleaned: list[str] = []
    for line in lines:
        normalized = " ".join(line.split())
        if not cleaned or " ".join(cleaned[-1].split()) != normalized:
            cleaned.append(line)
    return cleaned


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        return "\n".join(_docx_lines(path))
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError("仅支持 PDF、DOCX 和 TXT 文件")


def preview_html(path: Path, filename: str) -> str:
    """Render a safe, local HTML preview for formats browsers cannot display."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        body = f"<pre>{escape(path.read_text(encoding='utf-8', errors='replace'))}</pre>"
    elif suffix == ".docx":
        document = Document(path)
        parts: list[str] = []
        for block in _docx_blocks(document):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style = (block.style.name if block.style else "").lower()
                tag = "h2" if "heading" in style or "标题" in style else "p"
                parts.append(f"<{tag}>{escape(text)}</{tag}>")
                continue
            seen_cells: set[str] = set()
            rows: list[str] = []
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_key = cell._tc.getroottree().getpath(cell._tc)
                    if cell_key in seen_cells:
                        continue
                    seen_cells.add(cell_key)
                    value = "<br>".join(escape(p.text.strip()) for p in cell.paragraphs if p.text.strip())
                    if value:
                        cells.append(f"<td>{value}</td>")
                if cells:
                    rows.append(f"<tr>{''.join(cells)}</tr>")
            if rows:
                parts.append(f"<table>{''.join(rows)}</table>")
        body = "".join(parts) or "<p>无法生成文档预览，请下载原文件查看。</p>"
    else:
        raise ValueError("该格式应由浏览器直接预览")

    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(filename)}</title>
<style>body{{margin:0;padding:34px 42px;color:#202923;background:#fff;font:14px/1.75 -apple-system,BlinkMacSystemFont,'Segoe UI','Noto Sans SC',sans-serif}}h2{{font-size:18px;margin:24px 0 8px}}p{{margin:7px 0;white-space:pre-wrap}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.75 inherit}}table{{width:100%;border-collapse:collapse;margin:14px 0}}td{{border:1px solid #dce3da;padding:9px;vertical-align:top}}</style>
</head><body>{body}</body></html>"""
