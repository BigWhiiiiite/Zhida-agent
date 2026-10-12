"""Bounded local document ingestion. Extraction is never a confirmed personal fact."""
from __future__ import annotations

import hashlib
import json
import os
import zipfile
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from sqlalchemy import text

from .messaging import enqueue_job
from .repository import new_id, require_row

PARSER_VERSION = "text-artifacts-v1"
MAX_BYTES = 25 * 1024 * 1024
MAX_TEXT = 200_000


def storage_root() -> Path:
    return Path(os.getenv("ZHIDA_V2_FILE_ROOT", str(Path(__file__).resolve().parents[2] / ".v2-runtime" / "files"))).resolve()


def extract_document(path: Path, mime: str) -> dict:
    blocks, warnings = [], []
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("document_exceeds_25mb")
    if mime == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("encrypted_pdf_requires_user_action")
        if len(reader.pages) > 50:
            raise ValueError("document_exceeds_50_pages")
        for index, page in enumerate(reader.pages):
            content = page.extract_text() or ""
            blocks.append({"source_path": f"page:{index + 1}", "text": content})
            if not content.strip():
                warnings.append(f"page:{index + 1}:no_text_requires_ocr")
    elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("document_uncompressed_size_exceeded")
        from docx import Document
        document = Document(path)
        for index, paragraph in enumerate(document.paragraphs):
            if paragraph.text.strip():
                blocks.append({"source_path": f"paragraph:{index + 1}", "text": paragraph.text})
        for index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                blocks.append({"source_path": f"table:{index + 1}/row:{row_index + 1}",
                               "text": " | ".join(cell.text for cell in row.cells)})
        # Floating drawings/text boxes are not claimed as extracted.
        warnings.append("docx_drawings_headers_textboxes_not_extracted")
    elif mime == "text/plain":
        blocks = [{"source_path": "text:1", "text": path.read_text(encoding="utf-8")}]
    else:
        raise ValueError("unsupported_document_type")
    if sum(len(block["text"]) for block in blocks) > MAX_TEXT:
        raise ValueError("document_extracted_text_limit_exceeded")
    return {"blocks": blocks, "warnings": warnings,
            "status": "needs_review" if any(b["text"].strip() for b in blocks) else "needs_ocr"}


class DocumentStore:
    def __init__(self, db, root: Path | None = None):
        self.db, self.root = db, (root or storage_root()).resolve()

    def ingest(self, user_id: str, path: Path) -> dict:
        """CLI-only import of an explicitly chosen file; no arbitrary-path API."""
        path = path.resolve(strict=True)
        if not path.is_file() or path.stat().st_size > MAX_BYTES:
            raise ValueError("invalid_document_size")
        suffix = path.suffix.casefold()
        mime = {".pdf": "application/pdf", ".txt": "text/plain",
                ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}.get(suffix)
        if mime is None:
            raise ValueError("only_pdf_docx_utf8_text_supported")
        content = path.read_bytes()
        if len(content) > MAX_BYTES:
            raise ValueError("invalid_document_size")
        checksum = hashlib.sha256(content).hexdigest()
        key = hashlib.sha256(user_id.encode()).hexdigest() + "/" + checksum + suffix
        with self.db.transaction(user_id) as cx:
            require_row(cx, "SELECT id FROM zhida.users WHERE id=:u", {"u": user_id})
            existing = cx.execute(text("SELECT id FROM zhida.files WHERE user_id=:u AND storage_key=:key"),
                                  {"u": user_id, "key": key}).scalar()
            file_id = existing or new_id()
            target = self.root / key
            target.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
            try:
                with target.open("xb") as stream:
                    stream.write(content)
                target.chmod(0o600)
            except FileExistsError:
                if hashlib.sha256(target.read_bytes()).hexdigest() != checksum:
                    raise ValueError("stored_document_hash_mismatch")
            if not existing:
                cx.execute(text("""INSERT INTO zhida.files(id,user_id,storage_key,content_hash,mime,size)
                    VALUES(:id,:u,:key,:hash,:mime,:size)"""),
                    {"id": file_id, "u": user_id, "key": key, "hash": checksum, "mime": mime, "size": len(content)})
            job = enqueue_job(cx, user_id=user_id, kind="document.parse", aggregate_id=file_id,
                              idempotency_key=f"parse:{file_id}:{PARSER_VERSION}", payload={"file_id": file_id})
        return {"file_id": file_id, "job_id": job.job_id}

    def parse(self, user_id: str, file_id: str) -> tuple[dict, object]:
        with self.db.transaction(user_id) as cx:
            file = require_row(cx, "SELECT * FROM zhida.files WHERE user_id=:u AND id=:id", {"u": user_id, "id": file_id})
            existing = cx.execute(text("""SELECT id,status FROM zhida.parsed_documents
                WHERE user_id=:u AND file_id=:id AND parser_version=:version"""),
                {"u": user_id, "id": file_id, "version": PARSER_VERSION}).mappings().first()
            if existing:
                return dict(existing), None
        target = (self.root / file["storage_key"]).resolve(strict=True)
        if not target.is_relative_to(self.root) or target.is_symlink():
            raise ValueError("document_storage_path_invalid")
        if target.stat().st_size != file["size"] or hashlib.sha256(target.read_bytes()).hexdigest() != file["content_hash"]:
            raise ValueError("document_changed_after_ingestion")
        result = extract_document(target, file["mime"])
        artifact_id = str(uuid5(NAMESPACE_URL, f"{user_id}:{file_id}:{PARSER_VERSION}"))

        def persist(cx):
            cx.execute(text("""INSERT INTO zhida.parsed_documents
                (id,user_id,file_id,parser_version,blocks,warnings,status)
                VALUES(:id,:u,:file,:version,CAST(:blocks AS jsonb),CAST(:warnings AS jsonb),:status)
                ON CONFLICT(user_id,file_id,parser_version) DO NOTHING"""),
                {"id": artifact_id, "u": user_id, "file": file_id, "version": PARSER_VERSION,
                 "blocks": json.dumps(result["blocks"], ensure_ascii=False),
                 "warnings": json.dumps(result["warnings"]), "status": result["status"]})
        return {"id": artifact_id, "status": result["status"], "block_count": len(result["blocks"])}, persist
